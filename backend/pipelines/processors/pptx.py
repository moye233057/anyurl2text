# -*- coding: utf-8 -*-
"""
PPTX 抽取。

设计要点：
- **不做无效图过滤**（小于 100x100、近似纯色、全黑等一律照收）：跳过判断只能靠猜，
  而漏掉的图往往正是流程图或截图。代价是 understand 可能对装饰性小图多花一点 token，
  由 MAX_EMBEDDED_IMAGES 上限兜住。
- 组合形状（GroupShape）递归展开——不递归会漏掉成组排版的正文。
"""
import re
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

import configs.settings as settings
from pipelines.processors import table_kind
from pipelines.processors.grid import GridBuilder, render_tsv
from pipelines.processors.schema import CoveredSlot, EmbeddedAsset, ExtractContext, ExtractResult

def _walk_shapes(shapes) -> list:
    """递归展开形状与组合形状，返回扁平列表。"""
    flat = []
    for shape in shapes:
        flat.append(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_walk_shapes(shape.shapes))
    return flat

def _shape_text(shape) -> str:
    """取形状里的文字；没有文本框时返回空串。"""
    if not getattr(shape, "has_text_frame", False):
        return ""
    return (shape.text_frame.text or "").strip()

def _merge_residue(grid, residues: dict) -> int:
    """
    把被覆盖格里的残留文字并回它的 origin，返回并入的处数。

    为什么会残留：合并区的非左上格是 DrawingML 的**续格**（`hMerge` / `vMerge`），产出方
    （WPS / LibreOffice / 手工改过的文件）有时把文字留在续格里。anydoc 直接丢弃这类文字，
    本仓的取向是「尽量抽出来」——并回 origin（空格分隔），既不丢字、也不改列位。
    去重是为了不重复：python-pptx 自己的 `merge()` 会把续格文字搬进 origin，两边都有的不并第二次。
    应用场景：`_table_text` 构造完网格后调用。
    输入：grid — `GridBuilder.finish()` 的网格；residues — {(行, 列): 文字}。
    输出：真正并入的处数。
    """
    merged = 0
    for (row, column), text in residues.items():
        if row >= len(grid) or column >= len(grid[row]):
            continue
        slot = grid[row][column]
        if not isinstance(slot, CoveredSlot):
            continue
        origin = grid[slot.origin_row][slot.origin_col]
        if text in origin.text:
            continue
        origin.text = f"{origin.text} {text}".strip()
        merged += 1
    return merged

def _table_text(shape) -> tuple:
    """
    把表格形状渲染成制表符分隔的行。

    功能：`row.cells` 在演示文稿里是一行一个 `a:tc`、顺序即列序；被覆盖的格由 `is_spanned` 判出
          交给格位网格，跨度只在合并区左上角取（python-pptx 明确说非 origin 上的跨度不可靠）。
    应用场景：extract_pptx 遇到表格形状时调用。
    输入：shape — 页上的形状。
    输出：(制表符分隔的多行文本, 命中的上界名列表, 表类型, 并入的残留文字处数)；不是表格时返回 ("", [], None, 0)。
    """
    if not getattr(shape, "has_table", False):
        return "", [], None, 0
    builder = GridBuilder()
    residues: dict = {}
    for row_index, row in enumerate(shape.table.rows):
        builder.next_row()
        for column_index, cell in enumerate(row.cells):
            # 格内空白压成单空格：一个格里有多段时，换行会把这一行切成两行，同表列数就对不上了
            text = re.sub(r"\s+", " ", cell.text or "").strip()
            if cell.is_spanned:
                builder.covered()
                if text:
                    residues[(row_index, column_index)] = text
                continue
            col_span = cell.span_width if cell.is_merge_origin else 1
            row_span = cell.span_height if cell.is_merge_origin else 1
            builder.place(text, col_span, row_span)
    grid = builder.finish()
    residue_count = _merge_residue(grid, residues)
    lines = render_tsv(grid, len(grid))
    return "\n".join(lines), list(builder.limit_hits), table_kind.classify(lines), residue_count

def extract_pptx(path: Path, ctx: ExtractContext) -> ExtractResult:
    """
    抽取 PPTX 每一页的文本、表格与图片信息。

    功能：逐页递归展开形状，按「文本 → 表格」顺序拼接；图片在需要时落盘为素材。
    应用场景：dispatch 分派 PPTX 时调用。
    输入：path — 文件路径；ctx — 抽取上下文。
    输出：ExtractResult，content 按页分割。
    异常：不抛——打开失败时返回降级结果并说明原因。
    """
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        return ExtractResult(
            degraded=True,
            degrade_reason=f"PPTX 打开失败: {exc}",
            degrade_code="pptx_parse_failed",
            engines=["python-pptx"],
        )

    pages = []
    images_total = 0
    assets = []
    seen = set()
    hits = []
    hidden_slides = []
    table_kinds = []
    residues = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        # 隐藏幻灯片（`p:sld@show="0"`）照抽——隐藏只是放映时跳过，不是没有内容；
        # 这里只登记是哪几张，让调用方知道产出里有「界面上看不到的页」（用户 2026-10-04 拍板）
        if (slide.element.get("show") or "1") == "0":
            hidden_slides.append(slide_index)
        blocks = []
        for shape in _walk_shapes(slide.shapes):
            text = _shape_text(shape)
            if text:
                blocks.append(text)
            table, table_hits, kind, residue_count = _table_text(shape)
            if table:
                blocks.append(table)
                table_kinds.append(kind)
            residues = residues + residue_count
            hits.extend(table_hits)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images_total += 1
                if not ctx.prepare_assets:
                    continue
                # 按图片部件的 sha1 去重，不按 rId：python-pptx 在包级按内容复用同一个
                # image part，而 rId 只在页内唯一——同一张图出现在两页上时 rId 可以不同，
                # 按 rId 去重会漏。sha1 与「同一个来源部件」在这条路上是一一对应的。
                fingerprint = shape.image.sha1
                if fingerprint in seen:
                    continue
                # 上限只在新图真的存不下时才算命中：剩下的全是重复引用就不是命中
                if len(assets) >= settings.MAX_EMBEDDED_IMAGES:
                    if "max_embedded_images" not in hits:
                        hits.append("max_embedded_images")
                    continue
                # 先记来源再落盘：同一张图只试一次，转码失败也不回头重试 N 遍
                seen.add(fingerprint)
                position = len(assets) + 1
                target = ctx.work_dir / f"pptx_s{slide_index}_img{position}.png"
                try:
                    with Image.open(BytesIO(shape.image.blob)) as image:
                        image.convert("RGB").save(target, format="PNG")
                except Exception:
                    continue
                assets.append(
                    EmbeddedAsset(
                        kind="image",
                        label=f"第{slide_index}页图片{position}",
                        path=target,
                        order=position,
                    )
                )
        title = f"=== 幻灯片 {slide_index} ==="
        pages.append(title + ("\n" + "\n".join(blocks) if blocks else ""))

    notes = []
    if residues:
        notes.append(f"覆盖格里的残留文字已并入合并单元格（{residues} 处）")
    if hidden_slides:
        listed = "、".join(f"第 {index} 张" for index in hidden_slides)
        notes.append(f"隐藏内容照抽：{listed}幻灯片是隐藏的（show=\"0\"）")

    return ExtractResult(
        content="\n\n".join(pages),
        engines=["python-pptx"],
        notes=notes,
        table_kind=table_kinds[0] if len(table_kinds) == 1 else None,
        pages=len(presentation.slides),
        pages_used=len(presentation.slides),
        images=images_total,
        images_used=len(assets),
        assets=assets,
        limits_hit=hits,
    )
